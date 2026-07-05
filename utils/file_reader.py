"""제출물 파일 읽기 모듈.

디렉토리 내 파일 읽기뿐만 아니라, 단일 Excel(.xlsx), CSV, PDF, TXT, DOCX 파일 업로드 및 분석을 지원합니다.
Excel/CSV의 경우 여러 학생의 기록이 들어있는 테이블(Table) 형태로 처리합니다.
"""

import os
import csv
import logging
from pathlib import Path
from typing import Optional

logger = logging.getLogger(__name__)


def _read_txt(file_path: str) -> Optional[str]:
    """텍스트 파일을 읽습니다. UTF-8 → CP949 순서로 시도합니다."""
    for encoding in ("utf-8", "cp949"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
        except OSError as e:
            logger.error("파일 읽기 실패 (%s): %s", file_path, e)
            return None
    logger.warning("인코딩 감지 실패, errors='replace'로 읽기 시도: %s", file_path)
    try:
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            return f.read()
    except OSError as e:
        logger.error("파일 읽기 최종 실패 (%s): %s", file_path, e)
        return None


def _read_docx(file_path: str) -> Optional[str]:
    """python-docx를 사용하여 .docx 파일의 본문 텍스트를 추출합니다."""
    try:
        from docx import Document
    except ImportError:
        logger.error("python-docx가 설치되지 않았습니다. pip install python-docx")
        return None

    try:
        doc = Document(file_path)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)
    except Exception as e:
        logger.error("DOCX 읽기 실패 (%s): %s", file_path, e)
        return None


def _read_pdf(file_path: str) -> Optional[str]:
    """pypdf를 사용하여 PDF 파일의 텍스트를 추출합니다."""
    try:
        from pypdf import PdfReader
    except ImportError:
        logger.error("pypdf 패키지가 설치되지 않았습니다. pip install pypdf")
        return None
    try:
        reader = PdfReader(file_path)
        pages_text = []
        for page in reader.pages:
            t = page.extract_text()
            if t:
                pages_text.append(t)
        return "\n".join(pages_text)
    except Exception as e:
        logger.error("PDF 읽기 실패 (%s): %s", file_path, e)
        return None


def _parse_filename(filename: str) -> tuple[str, Optional[str]]:
    """파일 이름에서 학생 식별자와 도서명을 분리합니다.

    형식: identifier_booktitle.ext → (identifier, booktitle)
    그 외: filename → (filename, None)
    """
    stem = Path(filename).stem
    if "_" in stem:
        parts = stem.split("_", maxsplit=1)
        identifier = parts[0].strip()
        book_title = parts[1].strip() if parts[1].strip() else None
        return identifier, book_title
    return stem, None


# 매핑 후보 헤더 리스트 (학번/이름/도서명/본문을 서로 다른 컬럼으로 명확히 분리)
_ID_HEADERS = ["학번", "번호", "student_id", "studentid", "student no", "student_no", "id"]
_NAME_HEADERS = ["이름", "성명", "성함", "학생명", "학생", "작성자", "제출자", "name", "student"]
_BOOK_HEADERS = ["도서명", "책제목", "책이름", "도서", "책", "제목", "book", "title"]
_TEXT_HEADERS = ["내용", "본문", "제출물", "독후감", "감상문", "글", "에세이", "text", "content", "essay"]

_BLANK_TOKENS = {"", "nan", "none", "null", "n/a", "na", "-", "미상", "없음"}


def _clean_cell(val) -> str:
    """셀 값을 문자열로 정규화한다 (줄바꿈 제거, 공백 트림)."""
    if val is None:
        return ""
    return str(val).replace("\n", " ").replace("\r", " ").strip()


def _is_blank(val: str) -> bool:
    """NaN/None/공백 등 결측치성 토큰인지 판정한다."""
    return val.strip().lower() in _BLANK_TOKENS


def _find_header_row(rows: list[list[str]], max_scan: int = 5) -> int:
    """첫 N개 행 중 학번/이름/도서명/본문 키워드와 가장 많이 일치하는 행을 헤더 행으로 판단한다.

    일부 Excel 파일은 실제 헤더 위에 제목/안내 문구 행이 있어, 첫 행(rows[0])을
    그대로 헤더로 가정하면 학번·이름·도서명 컬럼을 통째로 놓치는 결함이 있었다.
    """
    all_keywords = _ID_HEADERS + _NAME_HEADERS + _BOOK_HEADERS + _TEXT_HEADERS
    best_idx = 0
    best_score = -1
    for r_idx in range(min(max_scan, len(rows))):
        row = rows[r_idx]
        score = 0
        for cell in row:
            cell_lower = _clean_cell(cell).lower()
            if not cell_lower:
                continue
            if any(h in cell_lower for h in all_keywords):
                score += 1
        if score > best_score:
            best_score = score
            best_idx = r_idx
    # 최소 2개 이상의 키워드가 매칭된 행이 없으면 기존 동작대로 첫 행을 헤더로 취급
    if best_score < 2:
        return 0
    return best_idx


def _classify_columns(header: list[str]) -> dict[str, int]:
    """헤더 컬럼들을 학번/이름/도서명/본문 카테고리로 분류한다.

    각 컬럼은 아직 배정되지 않은 카테고리 중 우선순위(학번 > 이름 > 도서명 > 본문)에서
    처음 일치하는 하나에만 배정되어, 서로 다른 실제 컬럼이 뒤섞이지 않도록 한다.
    """
    idx = {"id": -1, "name": -1, "book": -1, "text": -1}
    for i, col in enumerate(header):
        col_lower = _clean_cell(col).lower()
        if not col_lower:
            continue
        if idx["id"] == -1 and any(h in col_lower for h in _ID_HEADERS):
            idx["id"] = i
        elif idx["name"] == -1 and any(h in col_lower for h in _NAME_HEADERS):
            idx["name"] = i
        elif idx["book"] == -1 and any(h in col_lower for h in _BOOK_HEADERS):
            idx["book"] = i
        elif idx["text"] == -1 and any(h in col_lower for h in _TEXT_HEADERS):
            idx["text"] = i
    return idx


def _guess_text_column(rows: list[list[str]], header_row_idx: int, claimed: set[int], header_len: int) -> int:
    """본문(text) 헤더를 못 찾았을 때, 아직 배정되지 않은 컬럼 중 데이터가 가장 긴
    (=독후감 본문일 가능성이 가장 높은) 컬럼을 휴리스틱으로 선택한다."""
    candidates = [i for i in range(header_len) if i not in claimed]
    if not candidates:
        return -1
    lengths = {i: 0 for i in candidates}
    for row in rows[header_row_idx + 1: header_row_idx + 1 + 30]:  # 표본 30행으로 충분
        for i in candidates:
            if i < len(row):
                lengths[i] += len(_clean_cell(row[i]))
    return max(candidates, key=lambda i: lengths[i])


def _parse_tabular_rows(rows: list[list[str]], file_path: str) -> list[dict]:
    """Excel/CSV 행 데이터에서 학번, 이름, 도서명, 본문 컬럼을 자동 매핑하여 읽어들입니다."""
    if not rows:
        return []

    header_row_idx = _find_header_row(rows)
    header = [_clean_cell(cell) for cell in rows[header_row_idx]]

    idx = _classify_columns(header)

    claimed = {v for v in idx.values() if v != -1}
    if idx["text"] == -1:
        idx["text"] = _guess_text_column(rows, header_row_idx, claimed, len(header))
    if idx["name"] == -1 and idx["id"] == -1:
        # 학번/이름 헤더를 전혀 찾지 못한 경우 하위 호환을 위해 0번 컬럼을 이름으로 취급
        idx["name"] = 0 if 0 not in claimed | {idx["text"]} else -1

    if idx["book"] == -1:
        logger.warning(
            "표 형식 제출물에서 '도서명' 헤더를 찾지 못했습니다 (파일: %s, 헤더: %s). "
            "도서명 없이 진행합니다.",
            Path(file_path).name, header,
        )

    results = []

    # 데이터 행 추출 (헤더 행 다음 행부터)
    for r_idx in range(header_row_idx + 1, len(rows)):
        row = rows[r_idx]
        if not row:
            continue

        def _cell_at(col_idx: int) -> str:
            if col_idx == -1 or col_idx >= len(row):
                return ""
            return _clean_cell(row[col_idx])

        id_val = _cell_at(idx["id"])
        name_val = _cell_at(idx["name"])
        book_val = _cell_at(idx["book"])
        text_val = _cell_at(idx["text"])

        id_val = "" if _is_blank(id_val) else id_val
        name_val = "" if _is_blank(name_val) else name_val
        book_val = None if _is_blank(book_val) else book_val
        text_val = "" if _is_blank(text_val) else text_val

        # 학번_이름 복합키 구성 (학번이 있으면 결합, 없으면 이름만 사용)
        if id_val and name_val:
            student_val = f"{id_val}_{name_val}"
        elif id_val:
            student_val = id_val
        elif name_val:
            student_val = name_val
        else:
            student_val = f"학생_{r_idx}"

        if not text_val:
            continue

        results.append({
            "student": student_val,
            "student_id": id_val,
            "student_name": name_val or student_val,
            "book_title": book_val,
            "text": text_val,
            "file_path": file_path,
            "file_type": "csv_row" if file_path.endswith(".csv") else "xlsx_row",
        })

    logger.info("표 형식 제출물 파싱 완료: %s 에서 %d행 추출 (헤더 행: %d, 컬럼 매핑: %s)",
                Path(file_path).name, len(results), header_row_idx, idx)
    return results


def _read_csv_file(file_path: str) -> list[dict]:
    """CSV 파일을 읽어서 학생 리스트를 반환합니다."""
    rows = []
    # UTF-8, CP949 인코딩 시도
    for encoding in ("utf-8-sig", "cp949", "utf-8"):
        try:
            with open(file_path, "r", encoding=encoding) as f:
                reader = csv.reader(f)
                rows = list(reader)
            break
        except UnicodeDecodeError:
            continue
        except OSError as e:
            logger.error("CSV 읽기 실패: %s", e)
            return []
            
    return _parse_tabular_rows(rows, file_path)


def _read_xlsx_file(file_path: str) -> list[dict]:
    """Excel(.xlsx) 파일을 읽어서 학생 리스트를 반환합니다."""
    try:
        import openpyxl
    except ImportError:
        logger.error("openpyxl 패키지가 설치되지 않았습니다. pip install openpyxl")
        return []
        
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        rows = []
        for r in sheet.iter_rows(values_only=True):
            # 행 전체가 None인 경우 패스
            if any(cell is not None for cell in r):
                rows.append(list(r))
        return _parse_tabular_rows(rows, file_path)
    except Exception as e:
        logger.error("Excel 읽기 에러 (%s): %s", file_path, e)
        return []


def read_submissions(submissions_path_or_dir: str) -> list[dict]:
    """제출물 경로(디렉토리 또는 파일, 또는 세미콜론으로 구분된 다중 파일)에서 학생 제출물을 읽어 목록을 반환합니다.

    Args:
        submissions_path_or_dir: 제출물이 저장된 디렉토리 또는 단일 파일 경로, 혹은 세미콜론 구분 경로.

    Returns:
        제출물 정보 딕셔너리 목록.
    """
    if not submissions_path_or_dir:
        return []
        
    # 세미콜론(;)으로 연결된 다중 파일 경로 처리
    if ";" in submissions_path_or_dir:
        results = []
        paths = submissions_path_or_dir.split(";")
        for p_str in paths:
            p_str = p_str.strip()
            if not p_str:
                continue
            results.extend(read_submissions(p_str))
        return results

    path = Path(submissions_path_or_dir)
    if not path.exists():
        logger.error("제출물 경로가 존재하지 않습니다: %s", submissions_path_or_dir)
        return []

    # 1. 단일 파일인 경우 처리
    if path.is_file():
        ext = path.suffix.lower()
        if ext == ".csv":
            return _read_csv_file(str(path))
        elif ext in (".xlsx", ".xls"):
            return _read_xlsx_file(str(path))
        elif ext == ".pdf":
            text = _read_pdf(str(path))
            student, book_title = _parse_filename(path.name)
            if text and text.strip():
                return [{
                    "student": student,
                    "student_id": "",
                    "student_name": student,
                    "book_title": book_title,
                    "text": text,
                    "file_path": str(path),
                    "file_type": "pdf"
                }]
            return []
        elif ext == ".txt":
            text = _read_txt(str(path))
            student, book_title = _parse_filename(path.name)
            if text and text.strip():
                return [{
                    "student": student,
                    "student_id": "",
                    "student_name": student,
                    "book_title": book_title,
                    "text": text,
                    "file_path": str(path),
                    "file_type": "txt"
                }]
            return []
        elif ext == ".docx":
            text = _read_docx(str(path))
            student, book_title = _parse_filename(path.name)
            if text and text.strip():
                return [{
                    "student": student,
                    "student_id": "",
                    "student_name": student,
                    "book_title": book_title,
                    "text": text,
                    "file_path": str(path),
                    "file_type": "docx"
                }]
            return []
        else:
            logger.error("지원하지 않는 파일 형식입니다: %s", path.name)
            return []

    # 2. 디렉토리인 경우 처리
    results: list[dict] = []
    supported_extensions = {".txt", ".docx", ".pdf", ".xlsx", ".xls", ".csv"}

    files = sorted(path.iterdir())
    for file_path in files:
        if not file_path.is_file():
            continue
        ext = file_path.suffix.lower()
        if ext not in supported_extensions:
            continue

        if ext == ".csv":
            results.extend(_read_csv_file(str(file_path)))
        elif ext in (".xlsx", ".xls"):
            results.extend(_read_xlsx_file(str(file_path)))
        else:
            student, book_title = _parse_filename(file_path.name)

            if ext == ".txt":
                text = _read_txt(str(file_path))
                file_type = "txt"
            elif ext == ".docx":
                text = _read_docx(str(file_path))
                file_type = "docx"
            elif ext == ".pdf":
                text = _read_pdf(str(file_path))
                file_type = "pdf"
            else:
                continue

            if text is None or not text.strip():
                continue

            results.append({
                "student": student,
                "student_id": "",
                "student_name": student,
                "book_title": book_title,
                "text": text,
                "file_path": str(file_path),
                "file_type": file_type,
            })

    logger.info("총 %d개 제출물 읽기 완료 (디렉토리: %s)", len(results), submissions_path_or_dir)
    return results
